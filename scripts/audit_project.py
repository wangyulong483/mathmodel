# -*- coding: utf-8 -*-
"""Offline audit for a CUMCM paper workspace.

The audit deliberately does not judge whether a mathematical argument is correct.
It checks the reproducibility and submission hazards that can be detected from the
input files, Markdown source, generated DOCX, and model outputs.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

try:
    from docx import Document
except ImportError:  # pragma: no cover - dependency is optional for Markdown-only audits
    Document = None

from project_utils import (
    configured_path,
    json_dump,
    load_config,
    parse_markdown_heading,
    print_header,
    project_path,
    read_text,
)


REQUIRED_HEADINGS = ["摘要", "1 问题重述", "2 问题分析", "3 模型假设", "4 符号说明"]
RECOMMENDED_HEADINGS = ["5 模型建立与求解", "6 模型评价", "参考文献", "附录"]
BANNED_FORMULA_COMMANDS = [r"\tag", r"\left", r"\right", r"\bigl", r"\bigr"]


def add_issue(report: dict[str, Any], level: str, code: str, message: str) -> None:
    report["issues"].append({"level": level, "code": code, "message": message})


def extract_markdown_tables(lines: list[str]) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    i = 0
    while i + 1 < len(lines):
        header = lines[i].strip()
        separator = lines[i + 1].strip()
        if "|" not in header or not re.match(r"^\|?\s*:?-{3,}", separator):
            i += 1
            continue
        columns = len([part for part in header.strip("|").split("|")])
        rows = 2
        i += 2
        while i < len(lines) and "|" in lines[i] and lines[i].strip():
            rows += 1
            i += 1
        tables.append({"columns": columns, "rows": rows})
    return tables


def audit_markdown(path: Path, config: dict[str, Any]) -> dict[str, Any]:
    text = read_text(path)
    lines = text.splitlines()
    try:
        target = str(path.relative_to(project_path(".").resolve()))
    except ValueError:
        target = str(path)

    report: dict[str, Any] = {
        "ok": True,
        "target": target,
        "issues": [],
        "metrics": {},
    }
    if not text.strip():
        add_issue(report, "error", "paper-missing-or-empty", "论文 Markdown 不存在或为空。")
        report["ok"] = False
        return report

    headings = [parse_markdown_heading(line) for line in lines]
    headings = [item for item in headings if item]
    heading_texts = [item[1] for item in headings]
    report["metrics"]["heading_count"] = len(headings)
    report["metrics"]["headings"] = heading_texts

    for expected in REQUIRED_HEADINGS:
        if not any(item == expected or item.startswith(expected + " ") for item in heading_texts):
            add_issue(report, "error", "missing-heading", f"缺少必要章节：{expected}")
    for expected in RECOMMENDED_HEADINGS:
        if not any(item == expected or item.startswith(expected + " ") for item in heading_texts):
            add_issue(report, "warning", "missing-recommended-heading", f"建议补齐章节：{expected}")

    if any("目录" in item for item in heading_texts):
        add_issue(report, "error", "table-of-contents", "国赛论文默认不放目录，请删除目录章节。")
    if re.search(r"(?im)^\s*(toc|table-of-contents)\s*:", text):
        add_issue(report, "error", "toc-metadata", "YAML 中启用了目录，请关闭 toc。")

    abstract_match = re.search(
        r"(?ms)^#\s*摘要\s*$\n(.*?)(?=^#\s+|\Z)",
        text,
    )
    if abstract_match:
        abstract = abstract_match.group(1)
        abstract_chars = len(re.sub(r"[*_`$#\[\]()>|\\]", "", abstract).strip())
        report["metrics"]["abstract_chars"] = abstract_chars
        if abstract_chars > 650:
            add_issue(report, "warning", "abstract-long", "摘要文本较长，转换后可能超过摘要专用页。")
        if not re.search(r"关键词", abstract):
            add_issue(report, "warning", "keywords-missing", "摘要页未检测到“关键词”。")
    else:
        add_issue(report, "error", "abstract-missing", "缺少摘要章节。")

    keyword_lines = [line for line in lines if "关键词" in line]
    if keyword_lines and any("；" in line or ";" in line for line in keyword_lines):
        add_issue(report, "warning", "keyword-separator", "关键词建议使用空格分隔，避免使用分号。")

    tables = extract_markdown_tables(lines)
    report["metrics"]["table_count"] = len(tables)
    report["metrics"]["table_columns"] = [table["columns"] for table in tables]
    for index, table in enumerate(tables, start=1):
        if table["columns"] > 6:
            add_issue(
                report,
                "warning",
                "wide-table",
                f"第 {index} 个 Markdown 表格有 {table['columns']} 列，建议拆分为不超过 6 列的窄表。",
            )

    formula_blocks = len(re.findall(r"\$\$", text))
    report["metrics"]["display_formula_delimiters"] = formula_blocks
    if formula_blocks % 2:
        add_issue(report, "error", "unbalanced-formula", "显示公式 $$ 定界符数量为奇数。")
    for command in BANNED_FORMULA_COMMANDS:
        if command in text:
            add_issue(report, "warning", "formula-command", f"检测到公式命令 {command}，可能影响 Pandoc/Word 渲染。")

    image_paths = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    missing_images = []
    for image in image_paths:
        image_path = (path.parent / image).resolve()
        if not image_path.exists():
            missing_images.append(image)
    report["metrics"]["image_count"] = len(image_paths)
    if missing_images:
        add_issue(report, "error", "missing-image", f"找不到图片：{', '.join(missing_images)}")

    caption_numbers = [int(n) for n in re.findall(r"\*\*表\s*(\d+)", text)]
    figure_numbers = [int(n) for n in re.findall(r"!\[图\s*(\d+)", text)]
    report["metrics"]["table_caption_numbers"] = caption_numbers
    report["metrics"]["figure_caption_numbers"] = figure_numbers
    for label, numbers in [("表", caption_numbers), ("图", figure_numbers)]:
        if numbers and numbers != list(range(1, len(numbers) + 1)):
            add_issue(report, "warning", f"{label}-numbering", f"{label}号未从 1 开始连续编号：{numbers}")

    reference_split = re.split(r"(?m)^#\s*参考文献\s*$", text, maxsplit=1)
    body_text = reference_split[0]
    reference_text = reference_split[1] if len(reference_split) > 1 else ""
    references = re.findall(r"(?m)^\s*\[(\d+)\]\s+", reference_text)
    citations = [int(n) for n in re.findall(r"\[(\d+)\]", body_text)]
    report["metrics"]["reference_count"] = len(references)
    report["metrics"]["citation_numbers"] = sorted(set(citations))
    if len(references) < 5:
        add_issue(report, "warning", "few-references", "参考文献少于 5 条，国赛稿件通常需要更充分的可核实来源。")
    uncited = sorted(set(map(int, references)) - set(citations))
    if uncited:
        add_issue(report, "warning", "uncited-reference", f"存在正文未引用的参考文献编号：{uncited}")

    if "答案" in text or "作业答案" in text:
        add_issue(report, "warning", "training-language", "检测到“答案/作业答案”等训练痕迹，请改为正式论文表述。")
    if re.search(r"(?i)\bSHAP\b", text) and not re.search(r"(?i)shap", read_text(configured_path(config, "model_script", "output/model.py"))):
        add_issue(report, "warning", "unsupported-method", "论文提到 SHAP，但模型代码中未检测到对应实现。")

    report["ok"] = not any(issue["level"] == "error" for issue in report["issues"])
    return report


def audit_docx(path: Path, report: dict[str, Any]) -> None:
    if not path.exists():
        add_issue(report, "warning", "docx-not-built", "尚未生成 paper.docx，跳过 DOCX 审计。")
        return
    if Document is None:
        add_issue(report, "warning", "python-docx-missing", "未安装 python-docx，跳过 DOCX 审计。")
        return

    try:
        doc = Document(path)
        report["metrics"]["docx_tables"] = len(doc.tables)
        report["metrics"]["docx_paragraphs"] = len(doc.paragraphs)
        with zipfile.ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            report["metrics"]["docx_images"] = sum(
                name.startswith("word/media/") for name in archive.namelist()
            )
            if "<w:keepNext" in document_xml:
                add_issue(report, "warning", "docx-keepnext", "DOCX 仍含 keepNext，可能造成跨页大片空白。")
            if "<w:tblStyle" not in document_xml:
                add_issue(report, "warning", "docx-table-style", "DOCX 中未检测到表格样式引用。")
            if not any("footer" in name.lower() for name in archive.namelist()):
                add_issue(report, "warning", "docx-footer", "DOCX 未检测到页脚部件，请确认页码已生成。")
    except Exception as exc:
        add_issue(report, "error", "docx-read-failed", f"DOCX 审计失败：{exc}")


def render_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# MathModel 项目审计报告",
        "",
        f"- 总体状态：{'PASS' if report['ok'] else 'FAIL'}",
        f"- 审计目标：`{report['target']}`",
        "",
        "## 指标",
        "",
    ]
    for key, value in report["metrics"].items():
        lines.append(f"- `{key}`: {json.dumps(value, ensure_ascii=False)}")
    lines.extend(["", "## 问题", ""])
    if not report["issues"]:
        lines.append("- 未发现自动化审计问题。")
    else:
        for issue in report["issues"]:
            lines.append(f"- **{issue['level'].upper()}** `{issue['code']}`：{issue['message']}")
    lines.extend(
        [
            "",
            "说明：该报告只覆盖可由文件与结构判断的风险；数学推导、参数合理性、图表解释和参考文献真实性仍需人工或联网复核。",
        ]
    )
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="离线审计数学建模项目")
    parser.add_argument("--paper", default=None, help="论文 Markdown 路径，默认读取配置")
    parser.add_argument("--docx", default=None, help="DOCX 路径，默认读取配置")
    parser.add_argument("--json", action="store_true", help="只输出 JSON")
    args = parser.parse_args()

    config = load_config()
    paper = Path(args.paper).resolve() if args.paper else configured_path(config, "paper_markdown", "output/paper.md")
    docx = Path(args.docx).resolve() if args.docx else configured_path(config, "paper_docx", "output/paper.docx")

    report = audit_markdown(paper, config)
    audit_docx(docx, report)
    report["ok"] = not any(issue["level"] == "error" for issue in report["issues"])

    json_path = configured_path(config, "audit_json", "output/audit_report.json")
    markdown_path = configured_path(config, "audit_markdown", "output/audit_report.md")
    json_dump(json_path, report)
    markdown_path.write_text(render_markdown(report), encoding="utf-8")

    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print_header("MathModel 离线审计")
        print(render_markdown(report))
        print(f"报告已写入：{json_path}")
        print(f"报告已写入：{markdown_path}")
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
