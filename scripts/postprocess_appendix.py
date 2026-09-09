# -*- coding: utf-8 -*-
"""Post-process and validate the generated DOCX.

The reference document owns typography and the three-line table style. This
script only performs instance-level cleanup that Pandoc cannot express safely:
caption styles, autofit layout, and removal of accidental keepNext/keepLines.
"""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def process_document(docx_path: Path) -> tuple[list[str], list[str]]:
    doc = Document(docx_path)
    warnings: list[str] = []
    errors: list[str] = []

    try:
        caption_style = doc.styles["Table Caption"]
    except KeyError:
        try:
            caption_style = doc.styles["Caption"]
        except KeyError:
            caption_style = None

    fixed_captions = 0
    if caption_style is not None:
        for table in doc.tables:
            previous = table._tbl.getprevious()
            if previous is None or previous.tag != qn("w:p"):
                continue
            text = "".join(node.text or "" for node in previous.iter() if node.tag == qn("w:t"))
            if "表" not in text:
                continue
            for paragraph in doc.paragraphs:
                if paragraph._element is previous:
                    paragraph.style = caption_style
                    fixed_captions += 1
                    break
    else:
        warnings.append("未找到 Table Caption/Caption 样式。")

    for table in doc.tables:
        table.autofit = True
        table_properties = table._tbl.find(qn("w:tblPr"))
        if table_properties is None:
            table_properties = etree.SubElement(table._tbl, qn("w:tblPr"))

        for child in list(table_properties):
            if child.tag in {qn("w:tblW"), qn("w:tblLayout")}:
                table_properties.remove(child)
        table_width = etree.SubElement(table_properties, qn("w:tblW"))
        table_width.set(qn("w:w"), "5000")
        table_width.set(qn("w:type"), "pct")
        layout = etree.SubElement(table_properties, qn("w:tblLayout"))
        layout.set(qn("w:type"), "autofit")

        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                if cell_properties is None:
                    cell_properties = etree.SubElement(cell._tc, qn("w:tcPr"))
                for child in list(cell_properties):
                    if child.tag == qn("w:tcW"):
                        cell_properties.remove(child)

    removed = 0
    for paragraph in doc.paragraphs:
        properties = paragraph._element.find(qn("w:pPr"))
        if properties is None:
            continue
        for tag in ("keepNext", "keepLines"):
            for child in list(properties):
                if child.tag == qn(f"w:{tag}"):
                    properties.remove(child)
                    removed += 1

    doc.save(docx_path)

    with zipfile.ZipFile(docx_path) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        if "<w:keepNext" in document_xml:
            errors.append("document.xml 仍包含 keepNext。")
        if not any(name.lower().startswith("word/footer") for name in names):
            warnings.append("未检测到页脚部件，请确认页码已生成。")

    reloaded = Document(docx_path)
    for index, table in enumerate(reloaded.tables, start=1):
        table_properties = table._tbl.find(qn("w:tblPr"))
        table_style = table_properties.find(qn("w:tblStyle")) if table_properties is not None else None
        if table_style is None:
            warnings.append(f"表格 {index} 未显式引用 Table 样式。")
        elif table_style.get(qn("w:val")) != "Table":
            warnings.append(f"表格 {index} 使用了 {table_style.get(qn('w:val'))}，不是 Table。")

    print(f"[OK] 修正表标题 {fixed_captions} 个，清除 keepNext/keepLines {removed} 个")
    print(f"[OK] 表格 {len(reloaded.tables)} 个，图片 {sum('image' in rel.reltype for rel in reloaded.part.rels.values())} 个")
    return warnings, errors


def main() -> int:
    parser = argparse.ArgumentParser(description="后处理并验证数学建模论文 DOCX")
    parser.add_argument("--docx", default="output/paper.docx", help="DOCX 文件路径")
    args = parser.parse_args()
    path = Path(args.docx).resolve()
    if not path.exists():
        print(f"[ERROR] 找不到 DOCX：{path}", file=sys.stderr)
        return 1

    warnings, errors = process_document(path)
    for warning in warnings:
        print(f"[WARN] {warning}")
    for error in errors:
        print(f"[ERROR] {error}", file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
